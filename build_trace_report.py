from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output' / 'TRACE_BTech_Project_Report_2026_2027.docx'
OUT.parent.mkdir(exist_ok=True)
IMG = ROOT / 'frontend' / 'src' / 'assets' / 'hero.png'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)

styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.12
for s, size in [('Title', 20), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
    styles[s].font.name = 'Times New Roman'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    styles[s].font.size = Pt(size); styles[s].font.bold = True; styles[s].font.color.rgb = RGBColor(0,0,0)

def border(section):
    p = section._sectPr
    b = OxmlElement('w:pgBorders'); b.set(qn('w:offsetFrom'), 'page')
    for edge in ('top','left','bottom','right'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'12'); e.set(qn('w:space'),'24'); e.set(qn('w:color'),'000000'); b.append(e)
    p.append(b)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=tcMar.find(qn('w:'+m))
        if el is None: el=OxmlElement('w:'+m); tcMar.append(el)
        el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; set_cell_shading(c,'1F4E79'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c)
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
        if widths: c.width=Inches(widths[i])
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            if ridx%2: set_cell_shading(cells[i],'EAF2F8')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    doc.add_paragraph()
    return t

def center(text='', size=12, bold=False, space=5):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(space)
    r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=bold; return p

def para(text='', boldlead=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if boldlead:
        r=p.add_run(boldlead); r.bold=True
    p.add_run(text); return p

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.add_run(text); return p

def chapter(no,title):
    doc.add_page_break(); center('TRACE',14,True,10); center(f'Chapter {no}',15,True,4); center(title,18,True,18)

def add_footer(section):
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=f.add_run('Department of Computer Science and Engineering, BVDU College of Engineering, Pune | '); r.font.name='Times New Roman'; r.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

border(sec); add_footer(sec)
# Cover
for _ in range(3): doc.add_paragraph()
center('A PROJECT REPORT',19,True,7); center('ON',14,False,17); center('TRACE',24,True,5); center('Technology for Reporting Analysis and Crime Evaluation',14,True,26)
center('Submitted to',13,False,2); center('BHARATI VIDYAPEETH',17,True,1); center('DEEMED TO BE UNIVERSITY',14,True,1); center('COLLEGE OF ENGINEERING, PUNE, INDIA - 411043',13,True,25)
center('In Partial Fulfillment of the Requirement for the Award of',12,False,12); center('BACHELOR OF TECHNOLOGY IN',15,True,1); center('COMPUTER SCIENCE AND ENGINEERING',15,True,15); center('BY',13,True,8)
t=doc.add_table(rows=3,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
for row in t.rows:
    for c in row.cells: set_cell_margins(c,80,100,80,100)
for i in range(3):
    t.cell(i,0).text='Exam No: __________________'; t.cell(i,1).text='Name: __________________'; t.cell(i,2).text='PRN No: ______________'
doc.add_paragraph(); center('UNDER THE GUIDANCE OF',13,True,2); center('Prof. ______________________________',13,True,24)
center('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING',13,True,3); center('BHARATI VIDYAPEETH DEEMED TO BE UNIVERSITY',13,True,2); center('COLLEGE OF ENGINEERING, PUNE, INDIA - 411043',13,True,2); center('2026-2027',12,True,0)

# Title page
doc.add_page_break(); center('A PROJECT REPORT',19,True,7); center('ON',14,False,15); center('TRACE',24,True,4); center('Technology for Reporting Analysis and Crime Evaluation',14,True,30)
center('Submitted to',13); center('BHARATI VIDYAPEETH DEEMED TO BE UNIVERSITY',15,True,2); center('COLLEGE OF ENGINEERING, PUNE, INDIA - 411043',13,True,25)
center('BACHELOR OF TECHNOLOGY IN COMPUTER SCIENCE AND ENGINEERING',14,True,15); center('Project group and guide details are to be filled before submission.',11,False,15)
center('UNDER THE GUIDANCE OF',13,True,3); center('Prof. ______________________________',13,True,27); center('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING',13,True,2); center('2026-2027',12,True,0)

# Certificate
doc.add_page_break(); center('BHARATI VIDYAPEETH DEEMED TO BE UNIVERSITY',14,True,2); center('COLLEGE OF ENGINEERING, PUNE, INDIA - 411043',12,True,20); center('CERTIFICATE',18,True,25)
para('This is to certify that the project entitled “TRACE - Technology for Reporting Analysis and Crime Evaluation” submitted by the students named on the cover page is a bona fide record of work carried out in partial fulfillment of the requirements for the award of the Bachelor of Technology degree in Computer Science and Engineering at Bharati Vidyapeeth Deemed to be University College of Engineering, Pune, during the academic year 2026-2027.')
doc.add_paragraph(); doc.add_paragraph();
table(['Date','Project Guide','Head of Department','External Examiner'], [['____________','____________________________','____________________________','____________________________']], [1.1,1.8,1.8,1.8])

# Acknowledgement / abstract
doc.add_page_break(); center('ACKNOWLEDGEMENTS',18,True,25)
para('We express our sincere gratitude to the Principal of Bharati Vidyapeeth Deemed to be University College of Engineering, Pune, and to the Head of the Department of Computer Science and Engineering for providing the academic environment and facilities required for this project.')
para('We are especially thankful to our project guide for the technical direction, review, and encouragement provided throughout the design and implementation of TRACE. We also thank all faculty members, peers, and contributors who supported the work.')
doc.add_paragraph(); para('Group Members: ______________________________')
doc.add_page_break(); center('ABSTRACT',18,True,25)
para('TRACE (Technology for Reporting, Analysis and Crime Evaluation) is an AI-assisted, geospatial incident reporting and safety-intelligence platform designed as a B.Tech Computer Science and Engineering project. The system helps citizens submit location-aware incident reports and helps authorized moderators review, verify, and manage those reports. It combines a React web interface, a Node.js and Express API, PostgreSQL with PostGIS for spatial storage and querying, and a FastAPI service for natural-language incident classification and preliminary severity estimation.')
para('The platform supports viewport-oriented safety-map queries, category and severity filters, verification-aware public visibility, role-based access control, notifications, audit logging, community corroboration, analytics, and real-time updates through Socket.IO. Its design keeps personal reporter information and unverified allegations away from the public view. AI results are used as preliminary triage signals and are not treated as proof of an incident. The work demonstrates an integrated application of web engineering, GIS, relational databases, machine learning, security, and real-time systems.')
para('Keywords: ', 'TRACE, incident reporting, geospatial intelligence, PostGIS, machine learning, RBAC, safety map')

# Contents
doc.add_page_break(); center('CONTENTS',18,True,18)
toc=[('1','Introduction'),('2','Literature Survey'),('3','Proposed Methodology'),('4','Software Requirements Specification'),('5','Requirement Analysis'),('6','System Design'),('7','Project Planning'),('8','Implementation'),('9','Screenshots of Results'),('10','System Testing'),('11','Conclusion and Future Scope'),('12','References')]
table(['Chapter','Title'],toc,[1,5.5])
center('LIST OF FIGURES',16,True,12)
table(['Figure','Caption'], [('3.1','TRACE logical architecture'),('6.1','Data and request flow'),('9.1','TRACE frontend asset')],[1,5.5])

chapter(1,'Introduction')
doc.add_heading('1.1 Background',level=2); para('Community safety information is often fragmented, delayed, and difficult to interpret geographically. TRACE addresses this problem through a map-first web application that turns a citizen report into structured, privacy-aware safety information for appropriate users. The platform is a decision-support and awareness prototype; it does not replace emergency services, police reporting, or legal investigation.')
doc.add_heading('1.2 Problem Statement',level=2); para('A practical incident reporting platform must collect narrative and location data, classify it consistently, prevent unverified claims from becoming public, support spatial exploration without loading an entire database into the browser, and allow authorized staff to review decisions. These needs motivate an integrated design rather than isolated reporting, mapping, and analytics features.')
doc.add_heading('1.3 Objectives',level=2)
for x in ['Provide authenticated citizens with a simple location-aware incident reporting workflow.','Classify report narratives and estimate preliminary severity through an AI service with a safe fallback.','Store incident coordinates in PostGIS and return only relevant map data for the current viewport.','Enable moderators and administrators to verify, manage, and audit report status changes.','Present trends, hotspots, and notifications while protecting reporter identities and sensitive details.'] : bullet(x)
doc.add_heading('1.4 Scope',level=2); para('The implemented scope covers a web frontend, REST APIs and WebSocket events, a PostgreSQL/PostGIS schema, AI inference endpoints, role-based access control, incident moderation, notifications, corroborations, safety-map functions, and analytics routes. Operational integration with emergency services, official police systems, or automated dispatch is outside the prototype scope.')

chapter(2,'Literature Survey')
doc.add_heading('2.1 Geospatial Incident Systems',level=2); para('Modern incident systems commonly combine geographic coordinates with structured attributes to support proximity, bounding-box, and density queries. PostGIS provides spatial types and indexes suitable for storing point incidents and filtering a map viewport efficiently. This project uses that approach instead of treating latitude and longitude as unrelated numbers.')
doc.add_heading('2.2 Text Classification for Triage',level=2); para('Classical natural-language processing pipelines such as text cleaning, TF-IDF feature extraction, and logistic regression are lightweight and explainable starting points for short incident narratives. TRACE uses this style of model behind a FastAPI interface and records the category, confidence, severity estimate, model version, and processing information separately from the incident record.')
doc.add_heading('2.3 Moderation Privacy and Trust',level=2); para('Crowdsourced reporting requires safeguards against misinformation and disclosure of reporter identities. TRACE addresses this with role-based access control, moderation states, public visibility flags, audit logs, notifications, and community corroboration. Public users receive only verified and public incident records, while staff can use the review workflow to make controlled decisions.')

chapter(3,'Proposed Methodology')
doc.add_heading('3.1 End to End Workflow',level=2); para('The system follows a controlled pipeline: a citizen submits a report; the backend validates the input and requests AI analysis; the report and AI result are stored in PostgreSQL/PostGIS; moderators review it; verified public records become visible through viewport-based map queries; analytics and live events keep the appropriate users informed.')
table(['Layer','Components','Responsibility'], [('Presentation','React, Vite, Leaflet','Map, report forms, dashboards, role-aware views'),('Application','Node.js, Express, TypeScript','REST APIs, validation, RBAC, audit logging, sockets'),('Intelligence','Python, FastAPI, scikit-learn','Narrative classification and severity estimation'),('Data','PostgreSQL, PostGIS','Users, incidents, spatial indexes, analytics data')],[1.25,2.1,3.15])
doc.add_heading('3.2 TRACE Logical Architecture',level=2)
table(['Citizen or Staff Browser','API and Real Time Layer','Data and Intelligence'], [('React and Leaflet interface','Express REST API\nJWT and RBAC\nSocket.IO','PostgreSQL plus PostGIS\nFastAPI classification and severity service')],[2.1,2.1,2.3])
center('Figure 3.1: TRACE logical architecture',9,False,12)
doc.add_heading('3.3 Privacy First Visibility',level=2); para('On report creation, records are submitted with a non-public status. The backend’s incident retrieval logic exposes only verified public records to standard public views, while report owners and authorized staff receive the scope appropriate to their roles. Reporter details and moderation notes are sanitized for non-staff viewers.')

chapter(4,'Software Requirements Specification')
doc.add_heading('4.1 Functional Requirements',level=2)
table(['ID','Requirement'], [('FR01','Register and authenticate citizens, moderators, and administrators.'),('FR02','Create a report with narrative, category, coordinates, location area, and occurrence time.'),('FR03','Classify the narrative and estimate preliminary severity.'),('FR04','Search and filter map incidents using geographic bounds, category, severity, and date.'),('FR05','Allow staff to review, update status, set visibility, and record moderation information.'),('FR06','Provide analytics, notifications, audit events, and real-time status signals.')],[.65,5.7])
doc.add_heading('4.2 Non Functional Requirements',level=2)
table(['Attribute','Requirement'], [('Security','JWT authentication, Helmet, CORS configuration, input validation, rate limiting, and role checks.'),('Privacy','Public views must not expose unverified reports, reporter contact details, or internal moderation notes.'),('Performance','Map APIs use spatial bounds and indexes; list APIs paginate results.'),('Reliability','AI calls have timeouts and the backend provides a heuristic fallback for classification and severity.'),('Maintainability','Frontend, backend, AI service, and database scripts are separated by responsibility.')],[1.2,5.15])
doc.add_heading('4.3 Software Environment',level=2); table(['Component','Technology'], [('Frontend','React, Vite, TypeScript, Leaflet'),('Backend','Node.js, Express, TypeScript, Socket.IO'),('Database','PostgreSQL with PostGIS'),('AI service','Python, FastAPI, scikit-learn'),('Development','Node.js 18+, Python 3.10+, PostgreSQL 15+ with PostGIS')],[1.7,4.65])

chapter(5,'Requirement Analysis')
doc.add_heading('5.1 User Roles',level=2)
table(['Role','Primary Access'], [('Citizen','Submit reports, see own reports, view permitted verified public safety data.'),('Moderator','Review incoming reports, update moderation state, inspect analysis, and manage queue decisions.'),('Administrator','Administrative controls, user management, system analytics, and audit oversight.')],[1.2,5.15])
doc.add_heading('5.2 Data Requirements',level=2); para('The core schema includes users, incidents, incident media metadata, incident analysis records, notifications, audit logs, corroborations, hotspot snapshots, and risk-grid records. Each incident includes a WGS84 PostGIS POINT location together with latitude and longitude values for convenient API use. GiST indexes support spatial filters, while regular indexes support status, category, severity, user, and time filters.')
doc.add_heading('5.3 Constraints and Assumptions',level=2); para('The platform relies on a configured PostgreSQL/PostGIS instance and an available AI service for model-backed inference. If the AI service is unavailable, the backend uses a limited keyword-based fallback, which preserves system continuity but does not provide the same model confidence. AI output remains a preliminary aid and requires moderation before public visibility.')

chapter(6,'System Design')
doc.add_heading('6.1 Data and Request Flow',level=2)
table(['Step','Processing'], [('1','User authenticates and sends report data to the incident API.'),('2','Backend validates input, calls the AI service, and chooses user category or predicted category.'),('3','Backend stores incident geometry, status, and analysis record in a transaction.'),('4','Audit and notification services record post-commit events; staff receive real-time updates.'),('5','Moderator verifies or rejects; verified public records may be broadcast to the public map.'),('6','Map and analytics endpoints query authorized, filtered, and geographically relevant data.')],[.6,5.75])
center('Figure 6.1: Data and request flow',9,False,12)
doc.add_heading('6.2 API Design',level=2)
table(['Route family','Purpose'], [('/api/auth','Registration, login, and authenticated identity operations.'),('/api/incidents','Report creation, retrieval, pagination, filtering, and staff moderation.'),('/api/map','Map data, spatial operations, and safe route features.'),('/api/analytics','Trend and hotspot oriented analytics.'),('/api/notifications','User notification retrieval and updates.'),('/api/admin','Administrative management and audit-oriented operations.')],[1.7,4.65])
doc.add_heading('6.3 Security Design',level=2); para('The backend applies security headers, configured cross-origin handling, request parsing limits, authentication middleware, role-based authorization middleware, validation middleware, and rate limiting. Audit entries capture key actions such as report creation and status changes. The database uses foreign keys and enum types to help maintain integrity.')

chapter(7,'Project Planning')
doc.add_heading('7.1 Development Plan',level=2)
table(['Phase','Planned Deliverable'], [('1 Foundation','Repository structure, API and service health checks.'),('2 Data Layer','PostgreSQL/PostGIS schema, indexes, migrations, and demo seed data.'),('3 Identity','Registration, login, JWT, and RBAC.'),('4 Map and Reporting','Interactive map, viewport queries, report form, and tracking.'),('5 Intelligence and Review','AI inference, moderation queue, verification, notifications.'),('6 Analytics and Hardening','Hotspots, trends, real-time updates, tests, and documentation.')],[1.35,4.95])
doc.add_heading('7.2 Risk Management',level=2)
table(['Risk','Mitigation'], [('False or harmful public information','Verification workflow, public visibility controls, anonymization, and audit logging.'),('Slow map interaction at scale','PostGIS GiST index, bounds-based queries, result limits, clustering, and pagination.'),('AI service unavailability','Short service timeout and heuristic fallback.'),('Sensitive data exposure','RBAC, owner/staff checks, API response sanitization, and environment-based secrets.')],[2.2,4.1])

chapter(8,'Implementation')
doc.add_heading('8.1 Frontend Implementation',level=2); para('The frontend is a React and Vite application written in TypeScript. It contains dedicated components for the landing view, authentication, registration, safety map, safe-route panel, report modal, corroboration modal, citizen report history, moderator queue, and administrator center. A central API service coordinates calls to the backend.')
doc.add_heading('8.2 Backend Implementation',level=2); para('The Express application creates an HTTP server and initializes Socket.IO. It registers route groups for authentication, incidents, maps, safe routes, analytics, notifications, administration, and corroboration. The incident controller performs AI analysis before opening the database transaction, stores the incident and analysis record atomically, then emits controlled notification and real-time events after commit.')
doc.add_heading('8.3 AI Implementation',level=2); para('The FastAPI service exposes health, classification, severity, and combined analysis endpoints. Its stated model version is 1.0.0-tfidf-logreg. The backend service wrapper calls the AI service with a three-second timeout. When the service cannot return a suitable result, it maps selected keywords to a category and derives a conservative severity estimate.')
doc.add_heading('8.4 Database Implementation',level=2); para('The database schema enables PostGIS and UUID support, defines role, category, severity, status, and notification enums, and maintains automatic updated-at values through triggers. The incidents table contains a geometry Point column with SRID 4326 and a GiST index. A partial composite index accelerates common verified-public map queries.')

chapter(9,'Screenshots of Results')
doc.add_heading('9.1 TRACE Frontend Asset',level=2); para('The repository includes the following frontend visual asset used by the application. Final submission can replace or supplement this page with captured screenshots of the running login, map, report, moderation, and analytics views after deployment.')
if IMG.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(IMG),width=Inches(2.45))
center('Figure 9.1: TRACE frontend asset',9,False,10)
doc.add_heading('9.2 Demonstrable User Journey',level=2)
for x in ['Citizen signs in and opens the safety map.','Citizen submits a location-aware incident narrative.','TRACE stores AI category and preliminary severity with the report.','Moderator receives the queue update and verifies or rejects the report.','Verified public data appears on permitted map and analytics views.']: bullet(x)

chapter(10,'System Testing')
doc.add_heading('10.1 Test Strategy',level=2); para('Testing focuses on critical paths: authentication, input validation, role boundaries, report creation, AI fallback behavior, spatial filtering, moderation updates, privacy filtering, and notification or socket events. The backend repository includes an API test entry point, while the service interfaces allow component-level health verification.')
doc.add_heading('10.2 Test Cases and Expected Results',level=2)
table(['ID','Test Case','Condition','Expected Result'], [('T01','Health check','Backend and database configured','API returns service status and database connection state.'),('T02','Citizen report creation','Valid JWT and valid narrative/location','Report is stored as SUBMITTED with AI analysis.'),('T03','Public incident list','Unauthenticated or citizen public request','Only VERIFIED and public records are returned.'),('T04','Moderator status update','Moderator or admin changes status','Record, audit entry, notification, and appropriate events are created.'),('T05','AI fallback','AI service is unavailable','Backend uses heuristic category and severity without failing the report.'),('T06','Spatial query','Map bounds are supplied','PostGIS bounding-box condition limits returned data.')],[.45,1.35,1.65,2.5])
doc.add_heading('10.3 Acceptance Criteria',level=2); para('The prototype is ready for demonstration when an authenticated user can submit a report, the system preserves the report and AI results, authorized staff can change its status, public endpoints protect unverified data, and map or analytics queries return controlled results. End-to-end verification should be performed against a configured local deployment before final submission.')

chapter(11,'Conclusion and Future Scope')
doc.add_heading('11.1 Conclusion',level=2); para('TRACE integrates incident reporting, AI-assisted text analysis, geographic data management, moderation, privacy controls, analytics, and real-time communication into a coherent safety-intelligence prototype. The codebase demonstrates a practical separation of frontend, backend, AI, and data responsibilities. Its verification-aware design is important because classification and severity estimates are useful for triage but must not be represented as facts without review.')
doc.add_heading('11.2 Future Scope',level=2)
for x in ['Add model evaluation reports, dataset governance, and multilingual narrative support.','Introduce richer spatial-temporal risk models only after validating data quality and moderation outcomes.','Add secure media storage and evidence-review workflows with retention controls.','Support mobile-friendly progressive web application features and opt-in local alerts.','Provide deployment automation, monitoring dashboards, and load testing for larger datasets.']: bullet(x)

chapter(12,'References')
refs=['TRACE project source code and repository documentation, inspected for implementation details, 2026.','PostGIS Documentation. PostgreSQL spatial and geographic object support. https://postgis.net/documentation/','FastAPI Documentation. https://fastapi.tiangolo.com/','scikit-learn Documentation. https://scikit-learn.org/stable/','React Documentation. https://react.dev/','Socket.IO Documentation. https://socket.io/docs/']
for i,r in enumerate(refs,1): para(f'[{i}] {r}')
doc.add_page_break(); center('RESEARCH PUBLICATION',18,True,25); para('Attach the published research paper or conference material here after publication or acceptance. This page is intentionally retained to match the supplied BVDU project-report structure.')

for s in doc.sections: add_footer(s)
doc.core_properties.title='TRACE BTech Project Report 2026 2027'
doc.core_properties.subject='Project documentation for TRACE'
doc.core_properties.author='TRACE Project Team'
doc.save(OUT)
print(OUT)
